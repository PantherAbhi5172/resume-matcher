import os
import json
from pathlib import Path

from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel,Field
from pypdf import PdfReader
from docx import Document

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("GROQ_API_KEY is not set.")

client = Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"




class MatchResult(BaseModel):
    score: float
    strengths: list[str] = Field(default_factory=list)
    missing_required_skills: list[str] = Field(default_factory=list)
    missing_preferred_skills: list[str] = Field(default_factory=list)
    meets_experience_requirement: bool
    experience_summary: str
    experience_match: str = ""
    concerns: list[str] = Field(default_factory=list)
    verdict: str = ""
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class JOB_D(BaseModel):
    role: str | None = None
    required_skills: list[str] = []
    preffered_skills: list[str] = []
    minimum_exp: float | None = None
    educational_requirement: list[str] = []
    responsibilities: list[str] = []    

jobd_schema = JOB_D.model_json_schema()

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience: float | None = None

    skills: list[str] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []
    experiences: list[Experience] = []


resume_schema = Resume.model_json_schema()



def read_pdf(file_path):
    reader = PdfReader(file_path)

    text = ""

    for page in reader.pages:
        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


def read_docx(file_path):
    doc = Document(file_path)

    text = ""

    for para in doc.paragraphs:
        if para.text.strip():
            text += para.text + "\n"

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return read_pdf(file_path)

    elif ext == ".docx":
        return read_docx(file_path)

    else:
        raise ValueError(
            "Unsupported file format. Only PDF and DOCX are supported."
        )  

def parse_job_description(job_description):

    system_prompt = f"""
    You are an HR assistant.

    Analyze the job description and extract relevant information.

    IMPORTANT:
    Your response MUST be valid JSON.
    Return ONLY JSON.
    Do not return markdown.
    Do not return explanations.

    Use this schema:

    {jobd_schema}

    SKILL EXTRACTION RULES:
    1. required_skills must contain ONLY concise skill names.
    2. preffered_skills must contain ONLY concise skill names.
    3. Prefer short names, normally 1-4 words.
    4. Extract things such as:
    - Programming languages
    - Frameworks
    - Libraries
    - Databases
    - Cloud platforms
    - Cloud services
    - Development tools
    - Technologies
    - Clearly defined technical competencies 

    5. DO NOT put complete sentences in the skill lists.

    6. DO NOT put years of experience in the skill lists.

    7. DO NOT put job responsibilities in the skill lists.

    8. DO NOT put educational requirements in the skill lists.

    9. If a requirement contains multiple technologies,
        extract the technologies individually when appropriate.   


     Examples:

BAD:
"5+ years of experience working on applications of machine learning"

GOOD:
"Machine Learning"

BAD:
"Experience with ML libraries such as Keras, TensorFlow, XGBoost"

GOOD:
"Keras"
"TensorFlow"
"XGBoost"

BAD:
"Proficiency with cloud services/tools (S3, EC2, EMR, SageMaker)"

GOOD:
"AWS"
"S3"
"EC2"
"EMR"
"SageMaker"

BAD:
"Experience building automated, supportable, monitored processes"

GOOD:
"Process Automation"

Return only information supported by the job description.
Do not invent skills.   
    """

    user_prompt = f"""
    Return the job description analysis as JSON.

    Job Description:
    {job_description}
    """

    messages = [
        {
            "role": "system",
            "content": system_prompt
        },
        {
            "role": "user",
            "content": user_prompt
        }
    ]

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={
            "type": "json_object"
        },
        temperature=0,
        max_tokens=2000
    )

    answer = response.choices[0].message.content

    raw_json = json.loads(answer)

    return JOB_D(**raw_json)

def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume




def final_score(job, resume):

    match_schema = MatchResult.model_json_schema()

    prompt = f"""
    You are an experienced HR recruiter evaluating a candidate
    against a job description.

    Compare the candidate's resume with the job description carefully.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Return ONLY valid JSON matching this schema:

    {match_schema}

    Evaluate the candidate on the following points:

    1. Overall match score from 0 to 100.
    2. Identify the candidate's strongest relevant skills and experience.
    3. Identify missing REQUIRED skills.
    4. Identify missing PREFERRED skills.
    5. Determine whether the candidate meets the experience requirement.
    6. Return the experience evaluation using these fields:
        - meets_experience_requirement: true or false
        - experience_summary: a concise one-line explanation

    7. Mention important concerns or weaknesses that an HR recruiter
       should know.
    8. Give a final recruitment recommendation.

    Important rules:

    - Required skills are more important than preferred skills.
    - Do not penalize the candidate heavily for missing preferred skills.
    - Do not invent experience or skills that are not present in the resume.
    - Base the evaluation only on the provided resume and job description.
    - Keep the response concise.
    - The verdict should be one of:
      "SHORTLIST",
      "CONSIDER",
      or
      "LOW FIT".
    - Return ONLY JSON.
    """

    message = {
        "role": "user",
        "content": prompt
    }

    messages = [message]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format,
        temperature=0
    )

    data = json.loads(
        response.choices[0].message.content
    )

    return MatchResult(**data)